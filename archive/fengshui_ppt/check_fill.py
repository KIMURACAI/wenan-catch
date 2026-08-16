#!/usr/bin/env python3
"""Check fill percentage of PPTX slides."""
from pptx import Presentation
from pptx.util import Emu

prs = Presentation('e:\\wenan catch\\成慈高级奇门讲义_高密度版.pptx')
total = len(prs.slides)
total_fill = 0; sample = 0

for i, s in enumerate(prs.slides):
    max_y = 0; min_y = 999999
    for sh in s.shapes:
        if sh.top is not None:
            y = sh.top / 914400
            h = (sh.height or 0) / 914400
            bottom = y + h
            if bottom > max_y: max_y = bottom
            if y < min_y: min_y = y
    fill_pct = min(100, (max_y - min_y) / 7.5 * 100)
    if i > 0 and i < total - 1:  # skip cover and end
        total_fill += fill_pct
        sample += 1
    if fill_pct < 75:
        print(f'  S{i+1}: {fill_pct:.0f}% fill (bottom at {max_y:.1f}\")')

avg = total_fill / sample if sample else 0
print(f'\nAverage fill: {avg:.0f}% ({sample} slides)')
print(f'Total: {total} slides')
