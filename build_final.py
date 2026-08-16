"""
用你裁切好的图 + 重新排版文字 → 输出 PPT
"""
import os, json, math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Kami 色板 ──
PARCHMENT  = RGBColor(0xf5, 0xf4, 0xed)
BRAND      = RGBColor(0x1B, 0x36, 0x5D)
NEAR_BLACK = RGBColor(0x14, 0x14, 0x13)
DARK_WARM  = RGBColor(0x3d, 0x3d, 0x3a)
OLIVE      = RGBColor(0x50, 0x4e, 0x49)
STONE      = RGBColor(0x6b, 0x6a, 0x64)
BORDER     = RGBColor(0xe8, 0xe6, 0xdc)

HEADING_FONT = "Source Han Serif SC"
BODY_FONT    = "Microsoft YaHei"

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# 用户裁切好的图片
USER_IMG_DIR = r"e:\wenan catch\_user_images"
JSON_PATH = r"C:\Users\Kimura\Desktop\待处理\comic_page_models.json"
OUTPUT = r"e:\wenan catch\output_final.pptx"

# 图表标注字符（删除）
LABEL_CHARS = set(
    "乾坤震巽坎离艮兑子丑寅卯辰巳午未申酉戌亥"
    "甲乙丙丁戊己庚辛壬癸东西南北中"
    "山水风火雷泽天地日月龙虎凤龟雀武"
    "向坐朝案穴池湖江河海井路门木金土"
)

def is_label(text, font_size, text_type):
    t = text.strip()
    if not t: return True
    if t.isdigit() and len(t) <= 3: return True
    if t == "中国传统风水学现代应用经典": return True
    if len(t) == 1 and t in "。，、；：？！""''…—～·": return True
    if font_size < 17 and all(c in LABEL_CHARS for c in t):
        return True
    return False


# ═══════════════ helpers ═══════════════

def blank_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = PARCHMENT; bg.line.fill.background()
    return s


def tb(slide, text, left, top, width, height,
       size=14, color=DARK_WARM, align=PP_ALIGN.LEFT,
       font=None, bold=False, spacing=1.35):
    if not text: return None
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]; p.alignment = align
    p.line_spacing = Pt(size * spacing)
    run = p.add_run(); run.text = text
    run.font.name = font or BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def brand_bar(slide, left, top, height, weight=2.5):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Pt(weight), height)
    bar.fill.solid(); bar.fill.fore_color.rgb = BRAND; bar.line.fill.background()


def add_line(slide, left, top, width, color=BORDER, weight=0.5):
    ln = slide.shapes.add_connector(1, left, top, left + width, top)
    ln.line.color.rgb = color; ln.line.width = Pt(weight)


def add_page_num(slide, num):
    tb(slide, str(num), Inches(12.6), Inches(7.1), Inches(0.5), Inches(0.3),
       size=9, color=STONE, align=PP_ALIGN.RIGHT)


def fit_image(slide, path, left, top, max_w, max_h):
    if not os.path.exists(path): return None
    from PIL import Image
    with Image.open(path) as im: iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    nw, nh = int(iw * scale), int(ih * scale)
    cx = left + (max_w - nw) // 2
    cy = top + (max_h - nh) // 2
    return slide.shapes.add_picture(path, cx, cy, nw, nh)


def estimate_text_height(text, font_size, box_w_inches):
    """估算文本行数"""
    chars_per_line = max(1, (box_w_inches * 72) / (font_size * 0.8))
    total_lines = 0
    for line in text.split('\n'):
        if not line.strip():
            total_lines += 1
        else:
            total_lines += max(1, math.ceil(len(line) / chars_per_line))
    return total_lines


# ═══════════════ 文字处理 ═══════════════

def clean_text_items(text_items):
    """过滤标注、合并相邻同Y的文字块"""
    # 过滤
    clean = []
    for item in text_items:
        t = item["text"].strip()
        fs = item.get("font_size", 10)
        tt = item.get("text_type", "")
        if is_label(t, fs, tt):
            continue
        clean.append(item)

    if not clean:
        return []

    # 按 Y 坐标分组（同一行）
    clean.sort(key=lambda x: x["bbox"][1])
    groups = []
    curr = [clean[0]]
    curr_y = clean[0]["bbox"][1]

    for item in clean[1:]:
        y = item["bbox"][1]
        if abs(y - curr_y) <= 18:  # 同一行
            curr.append(item)
        else:
            groups.append(curr)
            curr = [item]
            curr_y = y
    groups.append(curr)

    # 每组内按 X 排序，合并文字
    result = []
    for g in groups:
        g.sort(key=lambda x: x["bbox"][0])
        merged = "".join(t["text"] for t in g)
        avg_fs = sum(t.get("font_size", 10) for t in g) / len(g)
        avg_y = sum(t["bbox"][1] for t in g) / len(g)  # 保留 Y 坐标
        result.append({"text": merged, "font_size": avg_fs, "y": avg_y})

    return result


def organize_content(items):
    """分出标题和正文（标题必须在页面顶部 + 短 + 大字）"""
    real_headings = []
    body = []

    for item in items:
        txt = item["text"]
        fs = item["font_size"]
        y = item.get("y", 999)
        page_h = 842

        # 正文续句特征（不可能是标题）
        fragments = ["源、", "民父子", "力。", "究，", "气阳", "气阴",
                     "《葬书》", "的理论", "语。", "盛的一", "汇", "西、",
                     "的基础上", "会长久", "会长", "不会"]
        is_fragment = any(txt.startswith(f) for f in fragments)

        # 真正标题：页面顶部 + 字号 >= 18 + 短文本
        is_real_heading = (y < page_h * 0.3 and fs >= 18 and len(txt) <= 22 and not is_fragment)

        if is_real_heading:
            if txt not in real_headings:
                real_headings.append(txt)
        else:
            if len(txt) >= 2:
                body.append(txt)

    main_title = real_headings[0] if real_headings else ""
    subtitles = real_headings[1:] if len(real_headings) > 1 else []

    return main_title, subtitles, body


# ═══════════════ 主流程 ═══════════════

def get_user_images(page_num):
    """获取用户裁切的图片文件列表
    PPT 中 slide 1=封面, slide 2=第1页, slide 3=第2页...
    所以第 page_num 页的图片在 slide{page_num+1:03d}_img* """
    imgs = []
    slide_num = page_num + 1  # 封面占 slide 1
    for fname in sorted(os.listdir(USER_IMG_DIR)):
        if fname.startswith(f"slide{slide_num:03d}_img") and fname.endswith(('.png', '.jpg', '.jpeg')):
            imgs.append(os.path.join(USER_IMG_DIR, fname))
    return imgs


def main():
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        pages = json.load(f)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ═══ 封面 ═══
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BRAND; bg.line.fill.background()
    tb(s, "中国传统风水学\n现代应用经典",
       Inches(2.0), Inches(1.8), Inches(9.3), Inches(2.2),
       size=44, color=RGBColor(0xff, 0xff, 0xff), align=PP_ALIGN.CENTER,
       font=HEADING_FONT)
    add_line(s, Inches(5.5), Inches(4.3), Inches(2.3), PARCHMENT, 0.8)
    tb(s, "成慈先生 著",
       Inches(2.0), Inches(4.8), Inches(9.3), Inches(0.6),
       size=18, color=RGBColor(0xcc, 0xcc, 0xc0), align=PP_ALIGN.CENTER,
       font=BODY_FONT)

    # ═══ 内容页 ═══
    for page in pages:
        pn = page["page_num"]
        ptype = page.get("page_type", "")

        # 处理文字
        items = clean_text_items(page.get("text_items", []))
        main_title, subtitles, body = organize_content(items)

        # 获取用户裁切图
        user_imgs = get_user_images(pn)
        has_img = len(user_imgs) > 0

        s = blank_slide(prs)

        # ── 标题 ──
        content_top = Inches(0.3)
        if main_title:
            bar_h = Inches(0.5)
            brand_bar(s, Inches(0.35), content_top, bar_h)
            tb(s, main_title, Inches(0.55), content_top, Inches(9.0), bar_h,
               size=22, color=NEAR_BLACK, font=HEADING_FONT)
            body_top = content_top + Inches(0.75)
        else:
            body_top = content_top + Inches(0.4)

        # ── 底部分隔 + 页码 ──
        add_line(s, Inches(0.4), Inches(7.15), Inches(12.5), BORDER)
        add_page_num(s, pn)

        body_h = Inches(7.05) - body_top

        if not body and not subtitles:
            # 纯图
            if has_img:
                n = len(user_imgs)
                if n == 1:
                    fit_image(s, user_imgs[0], Inches(0.5), body_top, Inches(12.3), body_h)
                else:
                    # 多图垂直排列
                    gap = Inches(0.15)
                    each_h = (body_h - gap * (n - 1)) / n
                    for i, img_path in enumerate(user_imgs):
                        y = body_top + i * (each_h + gap)
                        fit_image(s, img_path, Inches(0.5), y, Inches(12.3), each_h)
            continue

        # ── 构造文字行 ──
        lines = []
        sizes_ = []
        colors_ = []
        fonts_ = []

        # 副标题
        for st in subtitles:
            lines.append(st)
            sizes_.append(15)
            colors_.append(BRAND)
            fonts_.append(HEADING_FONT)

        # 正文：按关键词区分重点
        for i, txt in enumerate(body):
            keywords = ["风水", "方法", "注意", "重要", "关键", "解决", "必须",
                       "不宜", "适宜", "选择", "要求", "原则", "禁忌", "宜", "忌"]
            is_key = any(kw in txt for kw in keywords)
            lines.append(txt)
            if is_key:
                sizes_.append(14)
                colors_.append(BRAND)
            else:
                sizes_.append(12)
                colors_.append(DARK_WARM)
            fonts_.append(BODY_FONT)

        # ── 估算文字高度 → 确定图文比例 ──
        if has_img:
            n_img = len(user_imgs)
            total_chars = sum(len(l) for l in lines)

            # 根据文字量决定图文宽度比例
            if total_chars < 80:
                img_w = Inches(8.8); text_left = Inches(9.5); text_w = Inches(3.6)
            elif total_chars < 200:
                img_w = Inches(7.5); text_left = Inches(8.1); text_w = Inches(5.0)
            elif total_chars < 400:
                img_w = Inches(6.2); text_left = Inches(6.8); text_w = Inches(6.3)
            else:
                img_w = Inches(5.2); text_left = Inches(5.8); text_w = Inches(7.3)

            # 图片区
            if n_img == 1:
                fit_image(s, user_imgs[0], Inches(0.25), body_top, img_w, body_h)
            else:
                gap = Inches(0.1)
                each_h = int((body_h - gap * (n_img - 1)) / n_img)
                for i, img_path in enumerate(user_imgs):
                    y = body_top + i * (each_h + gap)
                    fit_image(s, img_path, Inches(0.25), y, img_w, each_h)

            # 文字区：智能缩放防溢出
            # 估算所需高度
            est_total_lines = sum(
                estimate_text_height(l, sizes_[j] if j < len(sizes_) else 12, text_w / 914400)
                for j, l in enumerate(lines)
            )
            est_h_inches = est_total_lines * 14 * 1.25 / 72  # 14pt line

            # 如果超出，等比缩小
            scale_factor = 1.0
            if est_h_inches > body_h / 914400 * 1.1:
                scale_factor = (body_h / 914400) / est_h_inches * 0.9
                new_sizes = [max(9, int(s * scale_factor)) for s in sizes_]
            else:
                new_sizes = sizes_

            # 写入文字
            box = s.shapes.add_textbox(text_left, body_top, text_w, body_h)
            tf = box.text_frame; tf.word_wrap = True
            tf.margin_left = Pt(4); tf.margin_right = Pt(4)
            tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)

            for j, line in enumerate(lines):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                sz = new_sizes[j] if j < len(new_sizes) else new_sizes[-1]
                p.line_spacing = Pt(sz * 1.22)
                p.space_after = Pt(2)
                run = p.add_run(); run.text = line
                run.font.name = fonts_[j] if j < len(fonts_) else BODY_FONT
                run.font.size = Pt(sz)
                run.font.color.rgb = colors_[j] if j < len(colors_) else DARK_WARM
        else:
            # 纯文字：双栏或单栏
            total_chars = sum(len(l) for l in lines)
            if total_chars > 500:
                mid = len(lines) // 2
                col_w = Inches(5.9); gap = Inches(0.3)
                for side, start_idx in [("L", 0), ("R", mid)]:
                    end_idx = mid if side == "L" else len(lines)
                    sub_lines = lines[start_idx:end_idx]
                    sub_sizes = sizes_[start_idx:end_idx]
                    sub_colors = colors_[start_idx:end_idx]
                    sub_fonts = fonts_[start_idx:end_idx]
                    left = Inches(0.5) if side == "L" else Inches(0.5) + col_w + gap

                    box = s.shapes.add_textbox(left, body_top, col_w, body_h)
                    tf = box.text_frame; tf.word_wrap = True
                    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
                    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)

                    for j, line in enumerate(sub_lines):
                        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                        sz = sub_sizes[j] if j < len(sub_sizes) else 12
                        p.line_spacing = Pt(sz * 1.25)
                        p.space_after = Pt(3)
                        run = p.add_run(); run.text = line
                        run.font.name = sub_fonts[j] if j < len(sub_fonts) else BODY_FONT
                        run.font.size = Pt(sz)
                        run.font.color.rgb = sub_colors[j] if j < len(sub_colors) else DARK_WARM

                # 中间分隔线
                mid_x = Inches(0.5) + col_w + gap / 2
                add_line(s, mid_x, body_top + Inches(0.3), Inches(0.01), BORDER)
            else:
                # 单栏满宽
                box = s.shapes.add_textbox(Inches(0.5), body_top, Inches(12.3), body_h)
                tf = box.text_frame; tf.word_wrap = True
                tf.margin_left = Pt(4); tf.margin_right = Pt(4)
                tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)

                for j, line in enumerate(lines):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    sz = sizes_[j] if j < len(sizes_) else 12
                    p.line_spacing = Pt(sz * 1.25)
                    p.space_after = Pt(3)
                    run = p.add_run(); run.text = line
                    run.font.name = fonts_[j] if j < len(fonts_) else BODY_FONT
                    run.font.size = Pt(sz)
                    run.font.color.rgb = colors_[j] if j < len(colors_) else DARK_WARM

    prs.save(OUTPUT)
    print(f"\nDone: {len(prs.slides)} slides -> {OUTPUT}")

if __name__ == "__main__":
    main()
