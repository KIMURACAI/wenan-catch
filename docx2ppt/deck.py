"""
Stage 3 · PPT 装配：封面 / 理论页 / 章节页 / 排盘页 / 断语页

Kami 设计 + 道教元素装饰：
- 16:9（13.333 × 7.5 in），羊皮纸底 #f5f4ed
- 每页：顶部细线 + 右上角轮转八卦符号 + 左下角太极
- 段落排版：CJK 全角按 1em 计宽、行距 1.3、段后距 6pt，防"一坨"文字并保证不溢出
- 排盘页：标题 + 参数卡(2×3芯片) + 九宫排盘图（带外框）+ 图注
"""
import math
import os
import sys
# 项目根目录入 sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from lxml import etree
from pptx.oxml.ns import qn

from docx2ppt.chart import draw_qimen_chart, add_chart_caption

# ---------- Kami token ----------
PARCH = RGBColor(0xF5, 0xF4, 0xED)
IVORY = RGBColor(0xFA, 0xF9, 0xF5)
SAND = RGBColor(0xE8, 0xE6, 0xDC)
INK = RGBColor(0x1B, 0x36, 0x5D)
INK_LIGHT = RGBColor(0x2D, 0x5A, 0x8A)
STONE = RGBColor(0x6B, 0x6A, 0x64)
BORDER = RGBColor(0xE8, 0xE6, 0xDC)
NEAR_BLACK = RGBColor(0x14, 0x14, 0x13)

# ---------- 道教符号 ----------
TRIGRAMS = ['☰', '☱', '☲', '☳', '☴', '☵', '☶', '☷']
TAIJI = '☯'
SYM_FONT = 'Segoe UI Symbol'

# ---------- 画布 ----------
SLIDE_W = 13.333
SLIDE_H = 7.5

# ---------- 版式常量 ----------
MARGIN = 0.5
TITLE_TOP = 0.32
TITLE_H = 0.58
TITLE_SIZE = 20
SUBTITLE_TOP = 1.02
SUBTITLE_SIZE = 11

BODY_LEFT = 0.6
BODY_TOP = 1.42
BODY_W = 12.13
BODY_H = 5.38          # 底边 6.80，留足安全边距

# 排盘页
PARAM_TOP = 1.28
CHIP_GAP = 0.14
CHIP_COLS = 3
CHIP_W = (12.13 - CHIP_GAP * (CHIP_COLS - 1)) / CHIP_COLS
CHIP_H = 0.54
CHART_SIZE = 4.5
CHART_LEFT = (SLIDE_W - CHART_SIZE) / 2
CHART_TOP = 2.7
CAPTION_TOP = 7.1

# 段落排版模型（保守，防溢出）
LINE_MULT = 1.3        # 行距倍数
SAFETY = 1.18          # 行高安全系数（CJK 实际行高 > 字号×倍数）
SPACING = 6            # 段后距 pt
BUFFER_IN = 0.22       # 框底安全缓冲（in），杜绝贴边
MIN_FONT = 9


# ---------- 基础工具 ----------
def _new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PARCH
    return slide


def _set_font_run(run, size, color=NEAR_BLACK, bold=False, font='微软雅黑'):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', font)


def _add_text(slide, x, y, w, h, text, size, color, bold=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font='微软雅黑',
              word_wrap=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _set_font_run(r, size, color, bold, font)
    return tb


def _add_rule(slide, x, y, w, color=BORDER, weight=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y),
                                    Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def _deco_symbol(slide, text, x, y, size, color, w=0.9, h=0.4,
                 align=PP_ALIGN.RIGHT, font=SYM_FONT):
    return _add_text(slide, x, y, w, h, text, size, color, align=align,
                     font=font, word_wrap=False)


# ---------- 段落排版（CJK 保守模型） ----------
def _para_lines(p, size, w_in):
    cpl = max(1, (w_in * 72) / (size * 1.0))
    return max(1, math.ceil(len(p) / cpl))


def _paras_height(paras, size, w_in):
    total = 0.0
    for p in paras:
        if not p or not p.strip():
            total += size * LINE_MULT * SAFETY * 12700
            continue
        total += (_para_lines(p, size, w_in) * size * LINE_MULT * SAFETY
                  * 12700 + SPACING * 12700)
    return total


def _base_size(paras):
    total = sum(len(p) for p in paras)
    if total < 100:
        return 17
    if total < 250:
        return 15
    if total < 500:
        return 14
    return 13


def paginate(paras, size, w_in, h_emu):
    pages = []
    cur = []
    cur_h = 0.0
    for p in paras:
        if not p or not p.strip():
            continue
        h = _paras_height([p], size, w_in)
        if cur and cur_h + h > h_emu:
            pages.append(cur)
            cur = [p]
            cur_h = h
        else:
            cur.append(p)
            cur_h += h
    if cur:
        pages.append(cur)
    return pages


def _set_paragraphs(textbox, paragraphs, base_override=None,
                    spacing_after=SPACING, color=NEAR_BLACK):
    """替换文本框内容为段落列表，字号自适应（保守模型防溢出）"""
    tf = textbox.text_frame
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    first_p = tf.paragraphs[0]
    for r in list(first_p.runs):
        r._r.getparent().remove(r._r)

    box_w_emu = textbox.width - Emu(91440)
    box_h_emu = textbox.height - Emu(91440) - Emu(914400 * BUFFER_IN)
    if box_h_emu <= 0:
        box_h_emu = Emu(914400 * 5)

    size = base_override if base_override is not None else _base_size(paragraphs)
    while size >= MIN_FONT:
        w_in = box_w_emu / 914400
        if _paras_height(paragraphs, size, w_in) <= box_h_emu:
            break
        size -= 0.5

    if paragraphs:
        first_run = first_p.add_run()
        first_run.text = paragraphs[0]
        _set_font_run(first_run, size, color)
        for t in paragraphs[1:]:
            p = tf.add_paragraph()
            r = p.add_run()
            r.text = t
            _set_font_run(r, size, color)
        for p in tf.paragraphs:
            p.line_spacing = LINE_MULT
            p.space_after = Pt(spacing_after)
    return size


def _body_textbox(slide, left=BODY_LEFT, top=BODY_TOP, w=BODY_W, h=BODY_H):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb


# ---------- 页面装饰 ----------
def _decorate(slide, pagenum, chart=False):
    """每页顶部细线 + 右上轮转八卦 + 左下太极"""
    _add_rule(slide, MARGIN, 0.22, SLIDE_W - 2 * MARGIN, BORDER, 0.75)
    _deco_symbol(slide, TRIGRAMS[(pagenum - 1) % 8], 12.42, 0.08, 13, INK_LIGHT)
    _deco_symbol(slide, TAIJI, 0.32, 6.86, 11, INK_LIGHT, align=PP_ALIGN.LEFT)


def _add_title(slide, text, sub="", size=TITLE_SIZE):
    tb = slide.shapes.add_textbox(Inches(MARGIN), Inches(TITLE_TOP),
                                  Inches(SLIDE_W - 2 * MARGIN), Inches(TITLE_H))
    tf = tb.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = text
    _set_font_run(r, size, INK, bold=True)
    _add_rule(slide, MARGIN, TITLE_TOP + TITLE_H + 0.08, 2.6, INK, 2)
    if sub:
        _add_text(slide, MARGIN, SUBTITLE_TOP, SLIDE_W - 2 * MARGIN, 0.3,
                  sub, SUBTITLE_SIZE, STONE)


# ---------- 封面 ----------
def add_cover(prs, deck):
    slide = _new_slide(prs)
    _add_rule(slide, MARGIN, 0.55, SLIDE_W - 2 * MARGIN, INK, 1.5)
    # 大太极
    _deco_symbol(slide, TAIJI, 5.92, 0.85, 80, INK, w=1.5, h=1.4,
                 align=PP_ALIGN.CENTER)
    # 主标题
    _add_text(slide, 0.8, 2.35, 11.7, 1.3, "奇门遁甲 · 预测案例", 44, INK,
              bold=True, align=PP_ALIGN.CENTER)
    # 副题
    _add_text(slide, 0.8, 3.7, 11.7, 0.5, "成慈老师 · 时家奇门实战预测 92–168 案例详解",
              16, STONE, align=PP_ALIGN.CENTER)
    # 八卦横排
    _add_text(slide, 1.0, 5.1, 11.3, 0.5, " ☰  ☱  ☲  ☳  ☴  ☵  ☶  ☷ ",
              18, INK_LIGHT, align=PP_ALIGN.CENTER, font=SYM_FONT)
    # 底部
    _add_text(slide, 0.8, 6.35, 11.7, 0.4, "三奇六仪 · 九星八门 · 天地人神",
              12, INK_LIGHT, align=PP_ALIGN.CENTER)
    _add_rule(slide, MARGIN, 6.85, SLIDE_W - 2 * MARGIN, INK, 1.5)


# ---------- 章节页 ----------
def add_divider(prs, text, sub=""):
    slide = _new_slide(prs)
    _decorate(slide, len(prs.slides))
    _add_rule(slide, (SLIDE_W - 1.6) / 2, 2.55, 1.6, INK, 2)
    # 八卦横排
    _add_text(slide, 1.5, 2.35, 10.3, 0.5, " ☰  ☱  ☲  ☳  ☴  ☵  ☶  ☷ ",
              16, INK_LIGHT, align=PP_ALIGN.CENTER, font=SYM_FONT)
    # 标题
    _add_text(slide, 1.0, 2.95, 11.3, 0.7, text, 26, INK, bold=True,
              align=PP_ALIGN.CENTER)
    if sub:
        _add_text(slide, 1.0, 3.85, 11.3, 0.4, sub, 13, STONE,
                  align=PP_ALIGN.CENTER)
    # 小太极
    _deco_symbol(slide, TAIJI, 6.32, 4.35, 18, INK_LIGHT, w=0.7, h=0.6,
                 align=PP_ALIGN.CENTER)


# ---------- 参数卡 ----------
def _strip_prefix(s, prefixes):
    s = s.strip()
    for pf in prefixes:
        if s.startswith(pf):
            return s[len(pf):].strip()
    return s


def _build_chips(params):
    chips = []
    if params.gongli:
        chips.append(("公历", _strip_prefix(params.gongli, ["公历：", "公历:"])))
    if params.nongli:
        chips.append(("农历", _strip_prefix(params.nongli, ["农历：", "农历:"])))
    if params.sizhu:
        chips.append(("四柱", _strip_prefix(params.sizhu, ["四柱：", "四柱:"])))
    if params.dunju:
        chips.append(("遁局", _strip_prefix(params.dunju, ["遁局：", "遁局:"])))
    if params.jieqi:
        chips.append(("节气", params.jieqi.strip()))
    if params.shixun:
        vals = [_strip_prefix(x, ["时旬：", "时旬:"]) for x in params.shixun]
        chips.append(("时旬 · 值符 · 值使", "；".join(vals)))
    return chips


def _fit_value_size(value, w_in, max_size=9.5):
    size = max_size
    while size > 7:
        if len(value) * (size * 1.0) / 72 <= w_in - 0.15:
            return size
        size -= 0.5
    return 7


def _add_chip(slide, left, top, w, h, label, value):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(left), Inches(top), Inches(w), Inches(h))
    shape.adjustments[0] = 0.12
    shape.fill.solid()
    shape.fill.fore_color.rgb = IVORY
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    _add_text(slide, left + 0.1, top + 0.03, w - 0.2, 0.2, label, 8, STONE,
              bold=True)
    vsize = _fit_value_size(value, w)
    _add_text(slide, left + 0.1, top + 0.24, w - 0.2, 0.26, value, vsize,
              NEAR_BLACK)


def _add_param_card(slide, params):
    chips = _build_chips(params)
    if not chips:
        return
    n = len(chips)
    rows = min(2, (n + CHIP_COLS - 1) // CHIP_COLS)
    for i, (label, value) in enumerate(chips):
        col = i % CHIP_COLS
        row = i // CHIP_COLS
        x = MARGIN + col * (CHIP_W + CHIP_GAP)
        y = PARAM_TOP + row * (CHIP_H + 0.1)
        _add_chip(slide, x, y, CHIP_W, CHIP_H, label, value)


# ---------- 排盘页 ----------
def _short_case(case):
    t = case.title
    for sep in ("：", ":"):
        if sep in t:
            return t.split(sep)[0]
    return t[:6]


def add_pan_page(prs, case, pagenum):
    slide = _new_slide(prs)
    _decorate(slide, pagenum, chart=True)
    _add_title(slide, case.title, sub=case.section)
    _add_param_card(slide, case.params)
    # 排盘图外框
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   Inches(CHART_LEFT - 0.12), Inches(CHART_TOP - 0.12),
                                   Inches(CHART_SIZE + 0.24), Inches(CHART_SIZE + 0.24))
    frame.fill.background()
    frame.line.color.rgb = INK_LIGHT
    frame.line.width = Pt(1)
    frame.shadow.inherit = False
    draw_qimen_chart(slide, case.cells, CHART_LEFT, CHART_TOP, CHART_SIZE,
                     zhishi_door=case.params.zhishi)
    # 图注
    p = case.params
    parts = []
    if p.dunju:
        parts.append(_strip_prefix(p.dunju, ["遁局：", "遁局:"]).split("日空")[0].strip())
    if p.zhifu:
        parts.append(f"值符{p.zhifu}")
    if p.zhishi:
        parts.append(f"值使{p.zhishi}门")
    if p.shixun_best:
        parts.append(f"{p.shixun_best}旬")
    if parts:
        add_chart_caption(slide, " · ".join(parts), 0, CAPTION_TOP)
    return slide


# ---------- 断语页 / 理论页 ----------
def add_verdict_pages(prs, case, pagenum):
    all_paras = list(case.narrative) + list(case.verdict)
    if not all_paras:
        return 0
    short = _short_case(case)
    size = _base_size(all_paras)
    pages = paginate(all_paras, size, BODY_W, int((BODY_H - BUFFER_IN) * 914400))
    for k, page_paras in enumerate(pages, 1):
        slide = _new_slide(prs)
        _decorate(slide, pagenum + k)
        if len(pages) == 1:
            title = f"{short} · 详解"
        else:
            title = f"{short} · 详解 {k}/{len(pages)}"
        _add_title(slide, title, sub=case.section)
        _set_paragraphs(_body_textbox(slide), page_paras, base_override=size)
    return len(pages)


def add_theory_pages(prs, section, pagenum):
    paras = section.intro
    if not paras:
        return 0
    size = _base_size(paras)
    pages = paginate(paras, size, BODY_W, int((BODY_H - BUFFER_IN) * 914400))
    for k, page_paras in enumerate(pages, 1):
        slide = _new_slide(prs)
        _decorate(slide, pagenum + k)
        if len(pages) == 1:
            title = section.title
        else:
            title = f"{section.title} · 续 {k}/{len(pages)}"
        _add_title(slide, title)
        _set_paragraphs(_body_textbox(slide), page_paras, base_override=size)
    return len(pages)


# ---------- 页码 ----------
def _add_page_num(prs, total):
    for i, slide in enumerate(prs.slides, 1):
        _add_text(slide, SLIDE_W - 1.0, SLIDE_H - 0.38, 0.7, 0.28,
                  f"{i} / {total}", 9, STONE, align=PP_ALIGN.RIGHT)


# ---------- 主流程 ----------
def build(deck, out_path):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    add_cover(prs, deck)

    page = 2
    for sec in deck.sections:
        if sec.kind == "theory":
            page += add_theory_pages(prs, sec, page)

    third = next((s for s in deck.sections if s.title.startswith("第三章")), None)
    if third:
        page += 1
        add_divider(prs, "第三章 · 时家奇门预测方法及范围")

    for sec in deck.sections:
        if sec.kind == "cases":
            page += 1
            add_divider(prs, sec.title)
            page += add_theory_pages(prs, sec, page)
            for case in sec.cases:
                page += 1
                add_pan_page(prs, case, page)
                page += add_verdict_pages(prs, case, page)

    _add_page_num(prs, len(prs.slides))
    prs.save(out_path)
    n = len(prs.slides)
    print(f"输出: {out_path} ({n} 页)")
    return n


if __name__ == "__main__":
    from docx2ppt.parse import parse_docx
    src = sys.argv[1] if len(sys.argv) > 1 else r"e:\wenan catch\预测案例92-168.docx"
    out = sys.argv[2] if len(sys.argv) > 2 else r"e:\wenan catch\81work\预测案例_输出.pptx"
    deck = parse_docx(src)
    build(deck, out)
