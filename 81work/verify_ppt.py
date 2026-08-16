"""自动校验 v2：正文段逐字 + 参数值(去标签) + 溢出预检"""
import sys, math
sys.path.insert(0, r"e:\wenan catch")
from docx2ppt.parse import parse_docx, RE_GONGLI, RE_NONGLI, RE_SIZHU, RE_DUNJU, RE_SHIXUN, RE_JIEQI
from pptx import Presentation

SRC = r"e:\wenan catch\预测案例92-168.docx"
PPTX = r"e:\wenan catch\81work\预测案例_输出.pptx"

deck = parse_docx(SRC)

content_paras = []      # 正文类：必须逐字出现
param_values = []       # 参数值：去标签后必须出现
param_lines_all = []
for s in deck.sections:
    content_paras += s.intro
    for c in s.cases:
        content_paras.append(c.title)
        content_paras += c.narrative
        content_paras += c.verdict
        for ln in c.params.lines:
            param_lines_all.append(ln)
            for rx in (RE_GONGLI, RE_NONGLI, RE_SIZHU, RE_DUNJU, RE_SHIXUN):
                if rx.match(ln.strip()):
                    v = ln.strip()
                    idx = v.find("：")
                    if idx >= 0:
                        param_values.append(v[idx + 1:].strip())
            if RE_JIEQI.match(ln.strip()):
                param_values.append(ln.strip())
        for c2 in c.params.shixun:
            pass

# 去重
def uniq(xs):
    s = set(); out = []
    for x in xs:
        if x not in s:
            s.add(x); out.append(x)
    return out

content_uniq = uniq(content_paras)
param_uniq = uniq(param_values)

prs = Presentation(PPTX)
runs = []
for slide in prs.slides:
    for sh in slide.shapes:
        if sh.has_text_frame:
            runs.append(sh.text_frame.text)
big = "\n".join(runs)

missing_content = [p for p in content_uniq if p not in big]
missing_param = [v for v in param_uniq if v not in big]
print(f"正文段(去重): {len(content_uniq)}  缺失: {len(missing_content)}")
for m in missing_content[:20]:
    print(f"  MISS: {m[:60]}")
print(f"参数值(去重): {len(param_uniq)}  缺失: {len(missing_param)}")
for m in missing_param[:15]:
    print(f"  MISS值: {m[:50]}")

# 溢出预检
over = []
for si, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        if sh.width and sh.height and sh.width / 914400 > 10:
            w_in = sh.width / 914400 - 0.1
            h_emu = sh.height - 91440
            txt = tf.text.strip()
            if not txt:
                continue
            size = 14
            for p in tf.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                        break
                break
            lines = sum(max(1, math.ceil(len(p.text) / max(1, (w_in * 72) / (size * 0.95))))
                        for p in tf.paragraphs if p.text)
            need = lines * size * 1.3 * 12700
            if need > h_emu:
                over.append((si, round(need / h_emu, 2), txt[:28]))
print(f"溢出风险: {len(over)}")
for o in over[:20]:
    print(f"  页{o[0]} {o[1]}x: {o[2]}")
print("\n校验完成")
