import sys, re
sys.path.insert(0, r"e:\wenan catch")
from pptx import Presentation
prs = Presentation(r"e:\wenan catch\81work\预测案例_输出.pptx")
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip().replace("\n", " ")
            # 过滤装饰符号 / 页码
            if len(t) <= 3 and not re.search(r"[\u4e00-\u9fa5]{2,}", t):
                continue
            if re.match(r"^\d+ / \d+$", t):
                continue
            texts.append(t)
    title = texts[0] if texts else "(无正文)"
    print(f"{i:3d} | {title[:42]:44s} | shapes={len(slide.shapes)}")
