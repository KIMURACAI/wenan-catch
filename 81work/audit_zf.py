"""验证值符宫高亮：每盘恰 1 个 2pt 粗边框宫格，且其内容含八神'值符'"""
import sys
sys.path.insert(0, r"e:\wenan catch")
from pptx import Presentation
from pptx.util import Pt

PPTX = r"e:\wenan catch\81work\预测案例_输出.pptx"
prs = Presentation(PPTX)

issues = []
checked = 0
for si, slide in enumerate(prs.slides, 1):
    if len(slide.shapes) < 80:
        continue
    checked += 1
    # 找 2pt 粗边框矩形（chart 区）
    thick = []
    shen_texts = []
    for sh in slide.shapes:
        if sh.shape_type == 1:
            try:
                if sh.line.width == Pt(2):
                    w = sh.width / 914400
                    if 1.3 < w < 1.6:
                        thick.append((sh.left / 914400, sh.top / 914400))
            except Exception:
                pass
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t in ("值符", "螣蛇", "太阴", "六合", "白虎", "玄武", "九地", "九天"):
                shen_texts.append((sh.left / 914400, sh.top / 914400, t))
    if len(thick) != 1:
        issues.append((si, f"粗边框宫格数={len(thick)}"))
        continue
    tx, ty = thick[0]
    # 该宫格矩形内的八神文字
    in_cell = [t for (x, y, t) in shen_texts
               if tx <= x + 0.4 <= tx + 1.4667 and ty <= y + 0.2 <= ty + 1.4667]
    if "值符" not in in_cell:
        issues.append((si, f"粗边框宫格内八神={in_cell}"))

print(f"排盘页 {checked} 张，异常 {len(issues)}")
for i in issues:
    print("  ", i)
print("OK" if not issues else "有异常")
