"""程序化几何审计：排盘页宫格内容位置 + 画布越界检查"""
import sys
sys.path.insert(0, r"e:\wenan catch")
from pptx import Presentation

PPTX = r"e:\wenan catch\81work\预测案例_输出.pptx"
SW, SH = 13.333, 7.5

prs = Presentation(PPTX)

# 1. 全盘越界检查
off = []
for si, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        try:
            l, t, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if l is None or w is None:
            continue
        l, t, w, h = l / 914400, t / 914400, w / 914400, h / 914400
        if l < -0.01 or t < -0.01 or l + w > SW + 0.01 or t + h > SH + 0.01:
            off.append((si, sh.shape_type, round(l, 2), round(t, 2),
                        round(l + w, 2), round(t + h, 2)))
print(f"越界元素: {len(off)}")
for o in off[:15]:
    print("  ", o)

# 2. 排盘页（shape数 90 的页）宫格内容定位检查
def find_pan_pages():
    out = []
    for si, slide in enumerate(prs.slides, 1):
        if len(slide.shapes) >= 80:
            out.append(si)
    return out

pan_pages = find_pan_pages()
print(f"\n排盘页: {pan_pages}")

# 宫格矩形 = 每排盘页的 9 个 RECTANGLE（chart 区域，大小接近 1.467in）
for si in pan_pages[:3]:
    slide = prs.slides[si - 1]
    rects = []
    for sh in slide.shapes:
        if sh.shape_type == 1:  # AUTO_SHAPE
            if sh.width and sh.height:
                w = sh.width / 914400
                if 1.3 < w < 1.6:
                    rects.append((sh.left / 914400, sh.top / 914400, w, sh.height / 914400))
    print(f"\n[排盘页 {si}] 宫格矩形数: {len(rects)}")
    # 找 textbox 内容并按位置归类到宫格
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip()
            texts.append((sh.left / 914400, sh.top / 914400, sh.width / 914400, t))
    # 每宫格(以 rect 中心定位)聚合其文字
    def cell_of(px, py, rects):
        best, bd = None, 1e9
        for (rx, ry, rw, rh) in rects:
            cx, cy = rx + rw / 2, ry + rh / 2
            d = ((px - cx) ** 2 + (py - cy) ** 2) ** 0.5
            if d < bd:
                bd, best = d, (cx, cy)
        return best if bd < 0.4 else None
    palace_of_cell = {}
    for (rx, ry, rw, rh) in rects:
        cx, cy = rx + rw / 2, ry + rh / 2
        # 洛书方位：row = round((cy-2.56)/1.4667), col = round((cx-4.4667)/1.4667)
        row = round((cy - 2.56) / 1.4667)
        col = round((cx - 4.4667) / 1.4667)
        palace_of_cell[(row, col)] = (rx, ry)
    print(f"  期望宫格原点: 左上(4.4667,2.56) 格距1.4667")
    for (px, py, pw, t) in texts:
        # 只关注图表区域 (2.4~7.1, 4.4~8.7)
        if 2.4 <= py <= 7.1 and 4.4 <= px <= 8.7:
            cc = cell_of(px + pw / 2, py + 0.05, rects)
            if cc:
                row = round((cc[1] - 2.56) / 1.4667)
                col = round((cc[0] - 4.4667) / 1.4667)
                print(f"    宫格(r{row},c{col}) 文本: {t[:22]}")
