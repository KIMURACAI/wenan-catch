"""提取 output_final.pptx 中所有文本，导出为 txt"""
from pptx import Presentation

prs = Presentation('output_final.pptx')

with open(r'e:\wenan catch\ppt_all_text.txt', 'w', encoding='utf-8') as f:
    for idx, slide in enumerate(prs.slides, 1):
        f.write(f"\n{'='*60}\n")
        f.write(f"第 {idx} 页\n")
        f.write(f"{'='*60}\n")
        for sh in slide.shapes:
            if sh.has_text_frame:
                text = sh.text_frame.text.strip()
                if text:
                    f.write(f"{text}\n")

print(f"提取完成 → ppt_all_text.txt")
