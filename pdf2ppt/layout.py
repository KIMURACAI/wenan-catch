"""
Stage 5 · 应用清理段落 + 图片排版（上文字、下图）
- 把 cleaned_paragraphs.json 的段落写回 PPT 主文本框（移到顶部全宽，字号自适应防溢出）
- 图片按原 PDF 排入底部全宽区域（保持宽高比不裁剪），多图网格
- Kami 风格 + 道教元素装饰（顶线、八卦、太极、分隔线）

输出: output_final.pptx
"""
import json, math
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR
from lxml import etree

SLIDE_W = 12188952
SLIDE_H = 6858000

# Kami 设计 token
BRAND     = RGBColor(0x1B, 0x36, 0x5D)   # 墨蓝 · 主强调色
INK_LIGHT = RGBColor(0x2D, 0x5A, 0x8A)   # 淡墨蓝（亮底上的符号色）
BORDER    = RGBColor(0xE8, 0xE6, 0xDC)   # 暖边框/分隔线
STONE     = RGBColor(0x6B, 0x6A, 0x64)   # 页码
COVER_SYM = RGBColor(0xAF, 0xBE, 0xD4)   # 深蓝封面上的符号色

# 道教元素：八卦（乾兑离震巽坎艮坤）+ 太极
TRIGRAMS = ['☰', '☱', '☲', '☳', '☴', '☵', '☶', '☷']
TAIJI = '☯'
SYM_FONT = 'Segoe UI Symbol'

# 装饰开关（默认关 = 与最终成品一致；需要道教/Kami 装饰时改为 True）
DECOR = False

# 上文字下图版式（16:9）
TEXT_AREA_LEFT = Emu(914400 * 0.4)
TEXT_AREA_TOP = Emu(914400 * 0.85)    # 标题下方
TEXT_AREA_W = Emu(914400 * 12.5)
TEXT_AREA_H = Emu(914400 * 2.95)      # 0.85 ~ 3.8in

IMG_AREA_LEFT = Emu(914400 * 0.4)
IMG_AREA_TOP = Emu(914400 * 4.0)
IMG_AREA_W = Emu(914400 * 12.5)
IMG_AREA_H = Emu(914400 * 3.1)        # 4.0 ~ 7.1in


def _is_big_textbox(sh):
    """正文主文本框候选：有文字、尺寸足够大（排除标题/页码/符号）"""
    try:
        return (sh.has_text_frame and len(sh.text_frame.text.strip()) >= 15
                and sh.width / 914400 > 3 and sh.height / 914400 > 2)
    except Exception:
        return False


def set_paragraphs(textbox, paragraphs, base_override=None):
    """替换文本框内容为段落列表，字号自适应"""
    tf = textbox.text_frame
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    first_p = tf.paragraphs[0]
    for r in list(first_p.runs):
        r._r.getparent().remove(r._r)

    box_w_emu = textbox.width - Emu(91440)
    box_h_emu = textbox.height - Emu(91440)
    if box_h_emu <= 0:
        box_h_emu = Emu(914400 * 5)

    if base_override is not None:
        base_size = base_override
    else:
        total_chars = sum(len(p) for p in paragraphs)
        if total_chars < 100: base_size = 17
        elif total_chars < 250: base_size = 15
        elif total_chars < 500: base_size = 14
        else: base_size = 13

    size = base_size
    while size >= 10:
        w_in = box_w_emu / 914400
        cpl = max(1, (w_in * 72) / (size * 0.95))
        total_lines = 0
        for p in paragraphs:
            if not p: total_lines += 1; continue
            total_lines += max(1, math.ceil(len(p) / cpl))
        if total_lines * size * 1.3 * 12700 <= box_h_emu:
            break
        size -= 0.5

    first_run = first_p.add_run()
    first_run.text = paragraphs[0] if paragraphs else ''
    _set_font(first_run, size)
    for p_text in paragraphs[1:]:
        p = tf.add_paragraph()
        run = p.add_run(); run.text = p_text
        _set_font(run, size)
    for p in tf.paragraphs:
        pPr = p._p.find(qn('a:pPr'))
        if pPr is None:
            pPr = etree.SubElement(p._p, qn('a:pPr'))
        lnSpc = pPr.find(qn('a:lnSpc'))
        if lnSpc is None:
            lnSpc = etree.SubElement(pPr, qn('a:lnSpc'))
        else:
            for c in list(lnSpc): lnSpc.remove(c)
        pct = etree.SubElement(lnSpc, qn('a:spcPct'))
        pct.set('val', '130000')
    return size


def _set_font(run, size_pt):
    run.font.name = '微软雅黑'
    run.font.size = Pt(size_pt)
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', '微软雅黑')


def fit_and_place(img, area_left, area_top, area_w, area_h):
    """保持宽高比缩放居中"""
    iw, ih = img.width, img.height
    if iw == 0 or ih == 0: return
    scale = min(area_w / iw, area_h / ih) * 0.97
    nw, nh = int(iw * scale), int(ih * scale)
    nx = area_left + (area_w - nw) // 2
    ny = area_top + (area_h - nh) // 2
    img.left, img.top, img.width, img.height = nx, ny, nw, nh


def place_images_bottom(slide, imgs):
    """上文字下图：图片排入底部全宽区域，多图网格（保持宽高比）"""
    if not imgs: return
    n = len(imgs)
    if n == 1:
        fit_and_place(imgs[0], IMG_AREA_LEFT, IMG_AREA_TOP, IMG_AREA_W, IMG_AREA_H)
        return
    if n <= 3: cols = n
    elif n <= 6: cols = 2
    else: cols = 3
    rows = (n + cols - 1) // cols
    gap = Emu(914400 * 0.15)
    cell_w = (IMG_AREA_W - gap * (cols - 1)) // cols
    cell_h = (IMG_AREA_H - gap * (rows - 1)) // rows
    for i, img in enumerate(imgs):
        col, row = i % cols, i // cols
        fit_and_place(img, IMG_AREA_LEFT + col * (cell_w + gap),
                      IMG_AREA_TOP + row * (cell_h + gap), cell_w, cell_h)


def _add_rule(slide, left_in, top_in, width_in, color, weight_pt=1.0):
    """画一条水平细线（Kami 分隔/品牌线）"""
    ln = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Emu(914400 * left_in), Emu(914400 * top_in),
        Emu(914400 * (left_in + width_in)), Emu(914400 * top_in))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight_pt)


def _add_sym(slide, text, left_in, top_in, size_pt, color,
             w_in=0.75, h_in=0.35, align=PP_ALIGN.RIGHT):
    """插入道教符号文字框（太极/八卦）"""
    box = slide.shapes.add_textbox(
        Emu(914400 * left_in), Emu(914400 * top_in),
        Emu(914400 * w_in), Emu(914400 * h_in))
    tf = box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = SYM_FONT
    r.font.size = Pt(size_pt)
    r.font.color.rgb = color


def decorate(slide, pagenum, has_imgs):
    """Kami 风格 + 道教元素装饰（克制的，不堆砌）"""
    _add_rule(slide, 0.4, 0.28, 12.5, BORDER, 0.75)
    _add_sym(slide, TRIGRAMS[(pagenum - 1) % 8], 12.15, 0.18, 13, INK_LIGHT)
    if has_imgs:
        _add_rule(slide, 0.4, 3.95, 12.5, BORDER, 0.5)
        _add_sym(slide, TAIJI, 12.1, 3.86, 11, INK_LIGHT)


def decorate_cover(slide):
    """封面：深蓝底上加大太极 + 八卦横排"""
    _add_sym(slide, TAIJI, 6.17, 5.15, 80, COVER_SYM, w_in=1.0, h_in=1.0,
             align=PP_ALIGN.CENTER)
    _add_sym(slide, '☰ ☱ ☲ ☳ ☴ ☵ ☶ ☷', 2.6, 6.55, 14, COVER_SYM,
             w_in=8.1, h_in=0.4, align=PP_ALIGN.CENTER)


def apply_and_layout(pptx_path, cleaned_json, output_path):
    """应用段落 + 重排版（上文字下图 + 装饰）"""
    cleaned = json.load(open(cleaned_json, 'r', encoding='utf-8'))
    prs = Presentation(pptx_path)

    applied = 0
    for snum, slide in enumerate(prs.slides, 1):
        # 1. 应用清理后的段落到顶部全宽文字区
        paras = cleaned.get(str(snum), [])
        if paras:
            candidates = [sh for sh in slide.shapes if _is_big_textbox(sh)]
            if candidates:
                tb = max(candidates, key=lambda s: s.width * s.height)
                # 删除其余大文字框（build_ppt 原始文字），避免重复/过期文字
                for sh in candidates:
                    if sh is not tb:
                        sh._element.getparent().remove(sh._element)
                tb.left, tb.top, tb.width, tb.height = \
                    TEXT_AREA_LEFT, TEXT_AREA_TOP, TEXT_AREA_W, TEXT_AREA_H
                set_paragraphs(tb, paras)
                applied += 1

        # 2. 图片排入底部区域（稍放大）
        imgs = [sh for sh in slide.shapes if sh.shape_type == 13]
        if imgs:
            place_images_bottom(slide, imgs)

        # 3. 装饰（默认关，与成品一致；需开启时把 DECOR 改为 True）
        if DECOR:
            if snum == 1:
                decorate_cover(slide)
            else:
                decorate(slide, snum, bool(imgs))

    prs.save(output_path)
    print(f"应用段落: {applied} 页")
    print(f"输出: {output_path} ({len(prs.slides)} 页)")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Stage 5: 应用+排版")
    parser.add_argument("pptx", help="output_raw.pptx")
    parser.add_argument("cleaned", help="cleaned_paragraphs.json")
    parser.add_argument("-o", "--output", default="output_final.pptx")
    args = parser.parse_args()
    apply_and_layout(args.pptx, args.cleaned, args.output)
