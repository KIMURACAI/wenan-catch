"""
Stage 2 · 初步生成 PPT
从 comic_page_models.json 生成 16:9 Kami 风格 PPT（左图右文）。
如果 _user_images/ 有用户裁切的图（slide001_img01.png 命名），优先使用。
输出: output_raw.pptx
"""
import os, json, math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PARCHMENT  = RGBColor(0xf5, 0xf4, 0xed)
BRAND      = RGBColor(0x1B, 0x36, 0x5D)
NEAR_BLACK = RGBColor(0x14, 0x14, 0x13)
DARK_WARM  = RGBColor(0x3d, 0x3d, 0x3a)
OLIVE      = RGBColor(0x50, 0x4e, 0x49)
STONE      = RGBColor(0x6b, 0x6a, 0x64)
BORDER     = RGBColor(0xe8, 0xe6, 0xdc)
WHITE      = RGBColor(0xff, 0xff, 0xff)

HEADING_FONT = "Source Han Serif SC"
BODY_FONT    = "Microsoft YaHei"
SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

try:
    from pdf2ppt.labels import LABEL_CHARS
except ImportError:
    from labels import LABEL_CHARS


def is_label(text, font_size, text_type):
    t = text.strip()
    if not t: return True
    if t.isdigit() and len(t) <= 3: return True
    if t == "中国传统风水学现代应用经典": return True
    if len(t) == 1 and t in "。，、；：？！\"\"''…—～·": return True
    if font_size < 17 and all(c in LABEL_CHARS for c in t):
        return True
    return False


def blank_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = PARCHMENT; bg.line.fill.background()
    return s


def tb(slide, text, left, top, width, height, size=14, color=DARK_WARM,
       align=PP_ALIGN.LEFT, font=None, bold=False, spacing=1.35):
    if not text: return None
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]; p.alignment = align
    p.line_spacing = Pt(size * spacing)
    run = p.add_run(); run.text = text
    run.font.name = font or BODY_FONT
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = color
    return box


def brand_bar(slide, left, top, height, weight=2.5):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Pt(weight), height)
    bar.fill.solid(); bar.fill.fore_color.rgb = BRAND; bar.line.fill.background()


def add_line(slide, left, top, width, color=BORDER, weight=0.5):
    ln = slide.shapes.add_connector(1, left, top, left + width, top)
    ln.line.color.rgb = color; ln.line.width = Pt(weight)


def add_page_num(slide, num):
    tb(slide, str(num), Inches(12.6), Inches(7.1), Inches(0.5), Inches(0.3),
       size=9, color=STONE, align=PP_ALIGN.RIGHT)


def fit_image(slide, path, left, top, max_w, max_h):
    if not os.path.exists(path): return None
    from PIL import Image
    with Image.open(path) as im: iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    nw, nh = int(iw * scale), int(ih * scale)
    cx = left + (max_w - nw) // 2
    cy = top + (max_h - nh) // 2
    return slide.shapes.add_picture(path, cx, cy, nw, nh)


def estimate_text_height(text, font_size, box_w_inches):
    chars_per_line = max(1, (box_w_inches * 72) / (font_size * 0.8))
    total_lines = 0
    for line in text.split('\n'):
        if not line.strip():
            total_lines += 1
        else:
            total_lines += max(1, math.ceil(len(line) / chars_per_line))
    return total_lines


def clean_text_items(text_items):
    """过滤标注、合并相邻同Y文字块"""
    clean = [item for item in text_items
             if not is_label(item["text"].strip(), item.get("font_size", 10), item.get("text_type", ""))]
    if not clean:
        return []
    clean.sort(key=lambda x: x["bbox"][1])
    groups = [[clean[0]]]
    curr_y = clean[0]["bbox"][1]
    for item in clean[1:]:
        y = item["bbox"][1]
        if abs(y - curr_y) <= 18:
            groups[-1].append(item)
        else:
            groups.append([item])
            curr_y = y
    result = []
    for g in groups:
        g.sort(key=lambda x: x["bbox"][0])
        merged = "".join(t["text"] for t in g)
        avg_fs = sum(t.get("font_size", 10) for t in g) / len(g)
        avg_y = sum(t["bbox"][1] for t in g) / len(g)
        result.append({"text": merged, "font_size": avg_fs, "y": avg_y})
    return result


def organize_content(items):
    """分出标题和正文"""
    real_headings, body = [], []
    fragments = ["源、", "民父子", "力。", "究，", "气阳", "气阴", "《葬书》",
                 "的理论", "语。", "盛的一", "汇", "西、", "的基础上", "会长久",
                 "会长", "不会"]
    for item in items:
        txt = item["text"]; fs = item["font_size"]; y = item.get("y", 999)
        is_fragment = any(txt.startswith(f) for f in fragments)
        is_real_heading = (y < 842 * 0.3 and fs >= 18 and len(txt) <= 22 and not is_fragment)
        if is_real_heading:
            if txt not in real_headings:
                real_headings.append(txt)
        else:
            if len(txt) >= 2:
                body.append(txt)
    main_title = real_headings[0] if real_headings else ""
    subtitles = real_headings[1:] if len(real_headings) > 1 else []
    return main_title, subtitles, body


def build_ppt(page_models_json, output_path, user_img_dir="_user_images",
              book_title="", book_author=""):
    """从 page_models 生成 PPT"""
    with open(page_models_json, "r", encoding="utf-8") as f:
        pages = json.load(f)

    os.makedirs(user_img_dir, exist_ok=True)
    user_imgs_by_slide = scan_user_images(user_img_dir)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ═══ 封面 ═══
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BRAND; bg.line.fill.background()
    title = book_title or os.path.splitext(os.path.basename(page_models_json))[0]
    tb(s, title, Inches(2.0), Inches(1.8), Inches(9.3), Inches(2.2),
       size=44, color=WHITE, align=PP_ALIGN.CENTER, font=HEADING_FONT)
    add_line(s, Inches(5.5), Inches(4.3), Inches(2.3), PARCHMENT, 0.8)
    if book_author:
        tb(s, f"{book_author} 著", Inches(2.0), Inches(4.8), Inches(9.3), Inches(0.6),
           size=18, color=RGBColor(0xcc, 0xcc, 0xc0), align=PP_ALIGN.CENTER, font=BODY_FONT)

    # ═══ 内容页 ═══
    for page in pages:
        pn = page["page_num"]
        items = clean_text_items(page.get("text_items", []))
        main_title, subtitles, body = organize_content(items)

        user_imgs = user_imgs_by_slide.get(pn + 1, [])
        has_img = len(user_imgs) > 0

        s = blank_slide(prs)
        content_top = Inches(0.3)
        if main_title:
            brand_bar(s, Inches(0.35), content_top, Inches(0.5))
            tb(s, main_title, Inches(0.55), content_top, Inches(9.0), Inches(0.5),
               size=22, color=NEAR_BLACK, font=HEADING_FONT)
            body_top = content_top + Inches(0.75)
        else:
            body_top = content_top + Inches(0.4)

        add_line(s, Inches(0.4), Inches(7.15), Inches(12.5), BORDER)
        add_page_num(s, pn)
        body_h = Inches(7.05) - body_top

        if not body and not subtitles:
            if has_img:
                _place_images(s, user_imgs, Inches(0.5), body_top, Inches(12.3), body_h)
            continue

        # 构造段落：(文本, 字号, 颜色, 字体)
        keywords = ["风水", "方法", "注意", "重要", "关键", "解决", "必须",
                    "不宜", "适宜", "选择", "要求", "原则", "禁忌", "宜", "忌"]
        paras = [(st, 15, BRAND, HEADING_FONT) for st in subtitles]
        for txt in body:
            is_key = any(kw in txt for kw in keywords)
            paras.append((txt, 14 if is_key else 12,
                          BRAND if is_key else DARK_WARM, BODY_FONT))

        if has_img:
            total_chars = sum(len(t) for t, *_ in paras)
            if total_chars < 80:
                img_w, text_left, text_w = Inches(8.8), Inches(9.5), Inches(3.6)
            elif total_chars < 200:
                img_w, text_left, text_w = Inches(7.5), Inches(8.1), Inches(5.0)
            elif total_chars < 400:
                img_w, text_left, text_w = Inches(6.2), Inches(6.8), Inches(6.3)
            else:
                img_w, text_left, text_w = Inches(5.2), Inches(5.8), Inches(7.3)

            _place_images(s, user_imgs, Inches(0.25), body_top, img_w, body_h)

            est_total_lines = sum(
                estimate_text_height(t, sz, text_w / 914400)
                for t, sz, _, _ in paras
            )
            est_h = est_total_lines * 14 * 1.25 / 72
            if est_h > body_h / 914400 * 1.1:
                sf = (body_h / 914400) / est_h * 0.9
                paras = [(t, max(9, int(sz * sf)), c, f) for t, sz, c, f in paras]

            box = s.shapes.add_textbox(text_left, body_top, text_w, body_h)
            _write_paragraphs(box, paras)
        else:
            total_chars = sum(len(t) for t, *_ in paras)
            if total_chars > 500:
                mid = len(paras) // 2
                col_w, gap = Inches(5.9), Inches(0.3)
                for side, start in [("L", 0), ("R", mid)]:
                    end = mid if side == "L" else len(paras)
                    left = Inches(0.5) if side == "L" else Inches(0.5) + col_w + gap
                    box = s.shapes.add_textbox(left, body_top, col_w, body_h)
                    _write_paragraphs(box, paras[start:end])
                mid_x = Inches(0.5) + col_w + gap / 2
                add_line(s, mid_x, body_top + Inches(0.3), Inches(0.01), BORDER)
            else:
                box = s.shapes.add_textbox(Inches(0.5), body_top, Inches(12.3), body_h)
                _write_paragraphs(box, paras)

    prs.save(output_path)
    print(f"\nDone: {len(prs.slides)} slides -> {output_path}")


def scan_user_images(user_img_dir):
    """扫描一次用户配图库，返回 {slide_num: [图片路径]}。
    slide 号从 1 起（封面占 slide 1），即 slide{NNN}_img* 的 NNN。"""
    by_slide = {}
    if not os.path.isdir(user_img_dir):
        return by_slide
    for fname in sorted(os.listdir(user_img_dir)):
        if not (fname.startswith('slide') and fname.endswith(('.png', '.jpg', '.jpeg'))):
            continue
        num = fname[5:8]
        if len(num) == 3 and num.isdigit() and fname[8:12] == '_img':
            by_slide.setdefault(int(num), []).append(os.path.join(user_img_dir, fname))
    return by_slide


def _place_images(slide, paths, left, top, max_w, max_h):
    n = len(paths)
    if n == 1:
        fit_image(slide, paths[0], left, top, max_w, max_h)
    else:
        gap = Inches(0.1)
        each_h = int((max_h - gap * (n - 1)) / n)
        for i, p in enumerate(paths):
            y = top + i * (each_h + gap)
            fit_image(slide, p, left, y, max_w, each_h)


def _write_paragraphs(box, paras):
    """paras: [(文本, 字号, 颜色, 字体)]"""
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    for j, (text, sz, color, font) in enumerate(paras):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.line_spacing = Pt(sz * 1.25); p.space_after = Pt(3)
        run = p.add_run(); run.text = text
        run.font.name = font
        run.font.size = Pt(sz)
        run.font.color.rgb = color


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Stage 2: 生成初步 PPT")
    parser.add_argument("json", help="comic_page_models.json")
    parser.add_argument("-o", "--output", default="output_raw.pptx")
    parser.add_argument("--imgdir", default="_user_images")
    parser.add_argument("--title", default="")
    parser.add_argument("--author", default="")
    args = parser.parse_args()
    build_ppt(args.json, args.output, args.imgdir, args.title, args.author)
