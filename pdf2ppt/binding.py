"""
Stage 3 · PPT ↔ PDF 页绑定
通过文本匹配建立每页 PPT 与 PDF 原始页的对应关系。
输出: slide_pdf_binding.json
"""
import json, re
from pptx import Presentation


def pdf_search_text(page):
    """PDF 页 → 供匹配的搜索文本（去掉页眉页码）"""
    return '\n'.join(
        t['text'].strip() for t in page.get('text_items', [])
        if t.get('text_type') != 'page_number' and t['text'].strip()
    )


def pdf_full_items(page):
    """PDF 页 → 带坐标的完整文字块列表"""
    return [{
        "text": t['text'].strip(),
        "bbox": t.get('bbox', []),
        "font_size": t.get('font_size', 0),
        "text_type": t.get('text_type', 'body'),
    } for t in page.get('text_items', []) if t['text'].strip()]


def pdf_entry_fields(page):
    """绑定条目中来自 PDF 页的三个字段"""
    return {
        "pdf_type": page.get('page_type', ''),
        "pdf_text": pdf_search_text(page),
        "pdf_text_with_pos": pdf_full_items(page),
    }


def build_binding(page_models_json, pptx_path, output_path):
    pages = json.load(open(page_models_json, 'r', encoding='utf-8'))
    prs = Presentation(pptx_path)

    def ppt_first_text(slide):
        for sh in slide.shapes:
            if sh.has_text_frame:
                for line in sh.text_frame.text.split('\n'):
                    line = line.strip()
                    if len(line) > 12 and not re.match(r'^第\d+页$', line):
                        return line
        return ''

    def ppt_all_text(slide):
        return '\n'.join(
            sh.text_frame.text.strip()
            for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()
        )

    # 预计算每页搜索文本，避免每页扫描时重复拼接 O(页数²)
    page_search_texts = [pdf_search_text(p) for p in pages]

    binding = {}
    for snum, slide in enumerate(prs.slides, 1):
        first = ppt_first_text(slide)
        pdf_page = None
        if first:
            key = first[:20]
            for p in pages:
                if key in page_search_texts[p['page_num'] - 1]:
                    pdf_page = p['page_num']
                    break

        entry = {"pdf_page": pdf_page, "ppt_first_text": first}
        if pdf_page:
            entry.update(pdf_entry_fields(pages[pdf_page - 1]))
        entry["ppt_text"] = ppt_all_text(slide)
        binding[f"slide_{snum}"] = entry

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(binding, f, ensure_ascii=False, indent=2)

    matched = sum(1 for e in binding.values() if e['pdf_page'])
    print(f"绑定完成: {matched}/{len(binding)} 页匹配")
    print(f"未匹配: {[k for k, v in binding.items() if not v['pdf_page']]}")
    return binding


def patch_manual(binding_path, manual_map, page_models_json):
    """手工指定未匹配页的对应 PDF 页"""
    binding = json.load(open(binding_path, 'r', encoding='utf-8'))
    pages = json.load(open(page_models_json, 'r', encoding='utf-8'))

    for key, pdf_page in manual_map.items():
        if key not in binding:
            continue
        e = binding[key]
        e['pdf_page'] = pdf_page
        e.update(pdf_entry_fields(pages[pdf_page - 1]))

    with open(binding_path, 'w', encoding='utf-8') as f:
        json.dump(binding, f, ensure_ascii=False, indent=2)
    matched = sum(1 for e in binding.values() if e['pdf_page'])
    print(f"手工补充后匹配: {matched}/{len(binding)}")
    return binding


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Stage 3: 建立 PPT↔PDF 绑定")
    parser.add_argument("json", help="comic_page_models.json")
    parser.add_argument("pptx", help="output_raw.pptx")
    parser.add_argument("-o", "--output", default="slide_pdf_binding.json")
    args = parser.parse_args()
    build_binding(args.json, args.pptx, args.output)
